import os
import docx
from pptx import Presentation

files = [
    "B\u00e0i 5. Qu\u1ea3n l\u00fd container v\u1edbi Docker Compose.pptx",
    "CI-CD Jenkins + Docker + GitHub .docx",
    "T\u1ea1o Cloud VPS b\u1eb1ng DigitalOcean.docx"
]

with open("output.txt", "w", encoding="utf-8") as f:
    for file in files:
        f.write(f"\n--- Reading {file} ---\n")
        try:
            if file.endswith(".docx"):
                doc = docx.Document(file)
                for para in doc.paragraphs:
                    if para.text.strip():
                        f.write(para.text + "\n")
            elif file.endswith(".pptx"):
                prs = Presentation(file)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            f.write(shape.text + "\n")
        except Exception as e:
            f.write(f"Error reading {file}: {e}\n")
