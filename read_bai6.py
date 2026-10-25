import os
from pptx import Presentation
import glob

file_path = glob.glob("*B*i 6*Jenkins.pptx")[0]

with open("output_bai6.txt", "w", encoding="utf-8") as f:
    f.write(f"\n--- Reading {file_path} ---\n")
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    f.write(shape.text + "\n")
    except Exception as e:
        f.write(f"Error reading: {e}\n")
